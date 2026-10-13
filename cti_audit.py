#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CTI REVERSE-ENGINEERING AUDIT  (cti_audit.py)
=============================================
Takes ANY table of metric columns (+ optional CTI summary / per-row CTI / model tables)
and produces an Excel audit workbook with live formulas:

  Summary · CTI_Fit · Weight_Spread · [Exact_Recovery] · Model_Consistency · Next_Steps

MODES
  A) Per-row CTI available (--cti-column)  -> exact recovery by OLS / NNLS (weights, R², residuals)
  B) Only CTI summary stats (--target or --target-from) -> distribution-matching fit of 4 formulations
     (simple weighted sum, equal weights, weighted geometric, affine) + weight-identifiability spread

INTEGRITY: mode B is a post-hoc fit to summary statistics. It cannot prove the original weights;
the Weight_Spread sheet shows how non-identifiable they are. Rows labelled SIMULATED stay flagged.

INSTALL   python -m pip install numpy pandas scipy openpyxl
EXAMPLES
  python cti_audit.py --data benchmark.xlsx --sheet Data_120 \
      --metrics "AI_Model_Accuracy (%)" "Batch_Failure_Rate (%)" "Production_Cost_Reduction (%)" \
                "Regulatory_Compliance_Score" "System_Resilience_Score" \
      --invert "Batch_Failure_Rate (%)" --target-from DADEH.xlsx --out CTI_audit.xlsx

  python cti_audit.py --data real.xlsx --metrics A B C D E --cti-column CTI --out audit.xlsx
  python cti_audit.py --data d.csv --metrics A B C --target "mean=.49,std=.08,min=.21,q25=.44,q50=.48,q75=.55,max=.67"
Optional: --models models.json (see DEFAULT_MODELS) ; --recalc (LibreOffice recalculation)
"""
import argparse, json, re, subprocess, sys, shutil
from pathlib import Path
import numpy as np, pandas as pd
from scipy.optimize import differential_evolution, nnls
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

KEYS = ["mean", "std", "min", "q25", "q50", "q75", "max"]
DEFAULT_MODELS = {   # example from the user's tables; override with --models
    "cti": [["Baseline ML", .41, .04, 0], ["XAI", .44, .05, 1.3], ["HITL", .58, .04, 5.3],
            ["Hybrid AI", .71, .03, 18.0], ["Hybrid + UQ", .83, .02, 93.4]],     # name, CTI, CI±, rank-stab %
    "pmi": [["Baseline ML", 42.3, 1.2, None, .229], ["XAI", 41.8, 1.1, .012, .224],
            ["HITL", 39.6, 1.0, .064, .214], ["Hybrid AI", 36.9, .9, .128, .201],
            ["Hybrid + UQ", 35.4, .8, .163, .193]]                               # name, PMI, CI±, reported red., energy-norm PMI
}

# ----------------------------------------------------------------------------- data
def read_table(path, sheet=None):
    p = Path(path)
    return pd.read_csv(p) if p.suffix.lower() == ".csv" else pd.read_excel(p, sheet_name=sheet or 0)

def parse_target_kv(s):
    d = {k: float(v) for k, v in (x.split("=") for x in s.split(","))}
    miss = [k for k in KEYS if k not in d]
    if miss: raise ValueError(f"--target missing keys: {miss}")
    return d

def parse_target_from(path):
    """Find 'mean 0.486 / std ... / 25% ... / max ...' text lines (pandas describe() pasted in cells)."""
    wb = load_workbook(path, data_only=True); lines = []
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            for c in row:
                if isinstance(c, str): lines.append(c.strip())
    start = next((i for i, l in enumerate(lines) if "composite trust index" in l.lower() or l.lower().startswith("cti")), 0)
    m = {"mean": "mean", "std": "std", "min": "min", "25%": "q25", "50%": "q50", "75%": "q75", "max": "max"}
    out = {}
    for l in lines[start:]:
        r = re.match(r"^(mean|std|min|25%|50%|75%|max)\s+(-?[\d.]+(?:[eE][-+]?\d+)?)$", l)
        if r and m[r.group(1)] not in out: out[m[r.group(1)]] = float(r.group(2))
    miss = [k for k in KEYS if k not in out]
    if miss: raise ValueError(f"Could not find CTI stats {miss} in {path}")
    return out

def normalise(X, invert, mode, ranges):
    Z = X.astype(float).copy()
    for c in Z.columns:
        if mode == "minmax":
            lo, hi = Z[c].min(), Z[c].max()
        else:
            lo, hi = ranges
        Z[c] = (Z[c] - lo) / (hi - lo)
        if c in invert: Z[c] = 1 - Z[c]
    return Z

def stats(c):
    c = np.asarray(c, float)
    return np.array([c.mean(), c.std(ddof=1), c.min(), *np.quantile(c, [.25, .5, .75]), c.max()])

# ----------------------------------------------------------------------------- fitting
def make_model(kind, M, L):
    k = M.shape[1]
    def pred(p):
        w = np.abs(p[:k]) + 1e-12; w = w / w.sum()
        if kind == "simple": return M @ w, w
        if kind == "geo":    return np.exp(L @ w), w
        if kind == "affine": return p[k] + p[k + 1] * (M @ w), w
        raise ValueError(kind)
    return pred

def fit(kind, M, tv, scale, seeds, maxiter):
    k = M.shape[1]; L = np.log(np.clip(M, 1e-3, 1)); pred = make_model(kind, M, L)
    bounds = [(0, 1)] * k + ([(-1, 1), (0.05, 3)] if kind == "affine" else [])
    loss = lambda p: float((((stats(pred(p)[0]) - tv) / scale) ** 2).sum())
    sols = []
    for s in range(seeds):
        r = differential_evolution(loss, bounds, seed=s, maxiter=maxiter, popsize=15, tol=1e-10, polish=True)
        c, w = pred(r.x); sols.append((r.fun, w, r.x[k:], stats(c)))
    sols.sort(key=lambda t: t[0])
    return sols

# ----------------------------------------------------------------------------- exact mode
def exact_recovery(M, y):
    A1 = np.c_[M, np.ones(len(M))]
    b_ols, *_ = np.linalg.lstsq(A1, y, rcond=None)
    w_nn, _ = nnls(M, y)
    def rep(pred): r = y - pred; return dict(R2=1 - (r ** 2).sum() / ((y - y.mean()) ** 2).sum(), maxabs=float(np.abs(r).max()))
    return dict(ols_w=b_ols[:-1], ols_b=b_ols[-1], ols=rep(A1 @ b_ols), nnls_w=w_nn, nnls=rep(M @ w_nn))

# ----------------------------------------------------------------------------- workbook
F = "Arial"; HDR = Font(name=F, bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor="1F3864")
IN = Font(name=F, color="0000FF", size=10); BK = Font(name=F, size=10); BO = Font(name=F, bold=True, size=10)
TT = Font(name=F, bold=True, size=14); NOTE = Font(name=F, italic=True, size=9, color="C00000")
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
RED = PatternFill("solid", fgColor="F8CBAD"); GRN = PatternFill("solid", fgColor="C6E0B4")
WR = Alignment(wrap_text=True, vertical="top")

def hdr(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.font = HDR; c.fill = HF; c.border = BD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
def put(ws, r, c, v, font=BK, fmt=None):
    x = ws.cell(r, c, v); x.font = font; x.border = BD
    if fmt: x.number_format = fmt
    return x
def L(i): return chr(64 + i)

def cf(ws, rng):
    a = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({a}="EXCEEDS",{a}="INCONSISTENT",{a}="CHECK")'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({a}="OK",{a}="CONSISTENT")'], fill=GRN))

def build(path, args, metrics, tgt, fits, spread, exact, models, n, simulated):
    wb = Workbook(); k = len(metrics)
    # ---- Summary (filled at the end)
    ws = wb.active; ws.title = "Summary"
    # ---- CTI_Fit
    if tgt is not None:
        ws = wb.create_sheet("CTI_Fit"); ws["A1"] = "CTI summary statistics: target vs fitted formulations"; ws["A1"].font = TT
        ws["A2"] = f"Normalisation: {args.norm}; inverted metrics: {args.invert or 'none'}; n = {n}."; ws["A2"].font = Font(name=F, italic=True, size=9)
        names = list(fits); hdr(ws, 4, ["Statistic", "Target"] + names + [f"Δ {x}" for x in names])
        for i, key in enumerate(KEYS, 5):
            put(ws, i, 1, key, BO); put(ws, i, 2, tgt[key], IN, "0.000")
            for j, nm in enumerate(names):
                put(ws, i, 3 + j, float(fits[nm][0][3][i - 5]), IN, "0.000")
                put(ws, i, 3 + len(names) + j, f"={L(3 + j)}{i}-$B{i}", BK, "+0.000;-0.000;0.000")
        r = 5 + len(KEYS)
        put(ws, r, 1, "Sum of squared Δ", BO); put(ws, r + 1, 1, "Max |Δ|", BO)
        for j in range(len(names)):
            col = L(3 + len(names) + j)
            put(ws, r, 3 + len(names) + j, f"=SUMSQ({col}5:{col}{r-1})", BK, "0.0000")
            put(ws, r + 1, 3 + len(names) + j, f"=MAX(MAX({col}5:{col}{r-1}),-MIN({col}5:{col}{r-1}))", BK, "0.000")
        r += 3; ws.cell(r, 1, "Fitted weights (sum = 1)").font = BO; hdr(ws, r + 1, ["Metric"] + names)
        for i, m in enumerate(metrics):
            put(ws, r + 2 + i, 1, m + (" (inverted)" if m in args.invert else ""))
            for j, nm in enumerate(names): put(ws, r + 2 + i, 2 + j, float(fits[nm][0][1][i]), IN, "0.000")
        rs = r + 2 + k; put(ws, rs, 1, "Sum of weights", BO)
        for j in range(len(names)): put(ws, rs, 2 + j, f"=SUM({L(2+j)}{r+2}:{L(2+j)}{rs-1})", BK, "0.000")
        if "affine" in fits:
            off, sc = fits["affine"][0][2]; put(ws, rs + 1, 1, "Affine offset / scale", BO)
            put(ws, rs + 1, 2, float(off), IN, "0.000"); put(ws, rs + 1, 3, float(sc), IN, "0.000")
        ws.cell(rs + 3, 1, "Fits to summary statistics are post-hoc. 'affine' has k+2 free parameters for 7 statistics: "
                           "a close match is expected and proves nothing about the true weights.").font = NOTE
        ws.column_dimensions["A"].width = 38
        for c in range(2, 2 + 2 * len(names) + 1): ws.column_dimensions[L(c)].width = 16
        # ---- Weight_Spread
        ws = wb.create_sheet("Weight_Spread"); ws["A1"] = "Weight identifiability: range of weights among near-optimal solutions"; ws["A1"].font = TT
        ws["A2"] = f"Solutions from independent optimiser seeds with loss ≤ max(2×best, best+{args.spread_tol})."; ws["A2"].font = Font(name=F, italic=True, size=9)
        r = 4
        for nm, (info) in spread.items():
            ws.cell(r, 1, nm).font = BO; hdr(ws, r + 1, ["Metric", "Min weight", "Max weight", "Range", "Mean weight", "Identifiable?"])
            for i, m in enumerate(metrics):
                put(ws, r + 2 + i, 1, m); put(ws, r + 2 + i, 2, float(info["lo"][i]), IN, "0.000"); put(ws, r + 2 + i, 3, float(info["hi"][i]), IN, "0.000")
                put(ws, r + 2 + i, 4, f"=C{r+2+i}-B{r+2+i}", BK, "0.000"); put(ws, r + 2 + i, 5, float(info["mean"][i]), IN, "0.000")
                put(ws, r + 2 + i, 6, f'=IF(D{r+2+i}<=0.05,"OK","CHECK")')
            cf(ws, f"F{r+2}:F{r+1+k}")
            put(ws, r + 2 + k, 1, "Solutions kept / seeds", BO); put(ws, r + 2 + k, 2, info["kept"], IN, "0"); put(ws, r + 2 + k, 3, info["seeds"], IN, "0")
            r += k + 5
        ws.cell(r, 1, "Range ≤ 0.05 (CHECK = weight not determined by the data) is a heuristic threshold, not a statistical test.").font = NOTE
        ws.column_dimensions["A"].width = 38
        for c in "BCDEF": ws.column_dimensions[c].width = 15
    # ---- Exact_Recovery
    if exact is not None:
        ws = wb.create_sheet("Exact_Recovery"); ws["A1"] = "Exact recovery from per-row CTI"; ws["A1"].font = TT
        hdr(ws, 3, ["Metric", "OLS weight (with intercept)", "NNLS weight (no intercept, ≥0)", "NNLS weight (normalised)"])
        for i, m in enumerate(metrics):
            put(ws, 4 + i, 1, m + (" (inverted)" if m in args.invert else "")); put(ws, 4 + i, 2, float(exact["ols_w"][i]), IN, "0.0000")
            put(ws, 4 + i, 3, float(exact["nnls_w"][i]), IN, "0.0000"); put(ws, 4 + i, 4, f"=IF(SUM($C$4:$C${3+k})=0,0,C{4+i}/SUM($C$4:$C${3+k}))", BK, "0.0000")
        r = 4 + k
        put(ws, r, 1, "Sum of weights", BO); put(ws, r, 2, f"=SUM(B4:B{r-1})", BK, "0.0000"); put(ws, r, 3, f"=SUM(C4:C{r-1})", BK, "0.0000")
        put(ws, r + 1, 1, "OLS intercept", BO); put(ws, r + 1, 2, float(exact["ols_b"]), IN, "0.0000")
        put(ws, r + 2, 1, "R² (OLS / NNLS)", BO); put(ws, r + 2, 2, exact["ols"]["R2"], IN, "0.000000"); put(ws, r + 2, 3, exact["nnls"]["R2"], IN, "0.000000")
        put(ws, r + 3, 1, "Max |residual| (OLS / NNLS)", BO); put(ws, r + 3, 2, exact["ols"]["maxabs"], IN, "0.000000"); put(ws, r + 3, 3, exact["nnls"]["maxabs"], IN, "0.000000")
        put(ws, r + 4, 1, "Verdict", BO); put(ws, r + 4, 2, f'=IF(B{r+2}>0.9999,"CTI is linear in these metrics (exact)",IF(B{r+2}>0.95,"Approximately linear","Not linear in these metrics: try other form"))')
        ws.column_dimensions["A"].width = 38
        for c in "BCD": ws.column_dimensions[c].width = 30
    # ---- Model_Consistency
    ws = wb.create_sheet("Model_Consistency"); ws["A1"] = "Consistency checks – model tables"; ws["A1"].font = TT
    mean_t = tgt["mean"] if tgt else None
    put(ws, 3, 1, "Distribution CTI max", BO); put(ws, 3, 2, tgt["max"] if tgt else float(exact["_max"]), IN, "0.000")
    put(ws, 4, 1, "Distribution CTI mean", BO); put(ws, 4, 2, tgt["mean"] if tgt else float(exact["_mean"]), IN, "0.000")
    put(ws, 5, 1, "Distribution CTI SD", BO); put(ws, 5, 2, tgt["std"] if tgt else float(exact["_std"]), IN, "0.000")
    put(ws, 6, 1, "n", BO); put(ws, 6, 2, n, IN, "0"); put(ws, 7, 1, "z for 95% CI", BO); put(ws, 7, 2, 1.96, IN, "0.00")
    hdr(ws, 9, ["Model", "CTI", "95% CI (±)", "Rank stability (%)", "Exceeds distribution max?", "vs distribution mean"])
    ct = models["cti"]; a, b = 10, 9 + len(ct)
    for i, (nm, v, ci, rs) in enumerate(ct, a):
        put(ws, i, 1, nm, BO); put(ws, i, 2, v, IN, "0.00"); put(ws, i, 3, ci, IN, "0.00"); put(ws, i, 4, rs, IN, "0.0")
        put(ws, i, 5, f'=IF(B{i}>$B$3,"EXCEEDS","OK")'); put(ws, i, 6, f'=IF(B{i}<$B$4,"Below","Above")')
    cf(ws, f"E{a}:E{b}")
    r = b + 2; ws.cell(r, 1, "Checks").font = BO; hdr(ws, r + 1, ["Check", "Value", "Reference", "Result"])
    put(ws, r + 2, 1, "Mean of model CTIs vs distribution mean"); put(ws, r + 2, 2, f"=AVERAGE(B{a}:B{b})", BK, "0.000"); put(ws, r + 2, 3, "=B4", BK, "0.000")
    put(ws, r + 2, 4, f'=IF(ABS(B{r+2}-C{r+2})<0.02,"CONSISTENT","INCONSISTENT")')
    put(ws, r + 3, 1, "Models above distribution max"); put(ws, r + 3, 2, f'=COUNTIF(E{a}:E{b},"EXCEEDS")', BK, "0"); put(ws, r + 3, 3, 0, IN, "0")
    put(ws, r + 3, 4, f'=IF(B{r+3}=C{r+3},"CONSISTENT","INCONSISTENT")')
    put(ws, r + 4, 1, "Expected CI half-width z·SD/√n vs mean reported"); put(ws, r + 4, 2, "=B7*B5/SQRT(B6)", BK, "0.000"); put(ws, r + 4, 3, f"=AVERAGE(C{a}:C{b})", BK, "0.000")
    put(ws, r + 4, 4, f'=IF(ABS(B{r+4}-C{r+4})<0.01,"CONSISTENT","INCONSISTENT")')
    cf(ws, f"D{r+2}:D{r+4}")
    pm = models.get("pmi") or []
    if pm:
        r = r + 7; ws.cell(r, 1, "PMI table").font = TT
        hdr(ws, r + 1, ["Model", "PMI", "95% CI (±)", "Reported reduction", "Recomputed reduction", "Energy-norm. PMI", "Implied energy", "Reduction check"])
        f0 = r + 2
        for i, (nm, v, ci, red, en) in enumerate(pm, f0):
            put(ws, i, 1, nm, BO); put(ws, i, 2, v, IN, "0.0"); put(ws, i, 3, ci, IN, "0.0")
            put(ws, i, 4, red if red is not None else "–", IN, "0.0%"); put(ws, i, 5, f"=(B${f0}-B{i})/B${f0}", BK, "0.0%")
            put(ws, i, 6, en, IN, "0.000"); put(ws, i, 7, f"=B{i}/F{i}", BK, "0.0")
            put(ws, i, 8, "n/a" if red is None else f'=IF(ABS(E{i}-D{i})<0.001,"OK","CHECK")')
        e = f0 + len(pm) - 1; cf(ws, f"H{f0}:H{e}")
        put(ws, e + 2, 1, "Implied-energy spread (% of mean)", BO); put(ws, e + 2, 7, f"=(MAX(G{f0}:G{e})-MIN(G{f0}:G{e}))/AVERAGE(G{f0}:G{e})", BK, "0.0%")
    ws.column_dimensions["A"].width = 46
    for c in "BCDEFGH": ws.column_dimensions[c].width = 19
    # ---- Next_Steps
    ws = wb.create_sheet("Next_Steps"); ws["A1"] = "What is needed to recover the CTI origin"; ws["A1"].font = TT
    hdr(ws, 3, ["#", "Item", "Why"])
    steps = [("Raw per-row CTI values (or CTI formula/code)", "With per-row CTI, weights follow from exact regression (Exact_Recovery sheet)."),
             ("Original (non-simulated) metric rows", "Simulated rows reproduce only mean/SD of the original."),
             ("n and CI method used for model CTIs", "Reported CIs are compared with z·SD/√n."),
             ("What each distribution/histogram represents", "Its range must contain every model CTI reported.")]
    for i, (a_, b_) in enumerate(steps, 4):
        put(ws, i, 1, i - 3); put(ws, i, 2, a_).alignment = WR; put(ws, i, 3, b_).alignment = WR; ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 5; ws.column_dimensions["B"].width = 55; ws.column_dimensions["C"].width = 75
    # ---- Summary
    ws = wb["Summary"]; ws["A1"] = "CTI reverse-engineering audit – summary"; ws["A1"].font = TT
    rows = [("Input data", f"{args.data}  (n = {n}; metrics: {', '.join(metrics)})"),
            ("Data status", "SIMULATED / not observed – results are illustrative only." if simulated else "Treated as supplied by the user; provenance not verified by this script."),
            ("Mode", "A: per-row CTI exact recovery" if exact is not None and tgt is None else
                     ("A+B: exact recovery and summary-statistics fit" if exact is not None else "B: fit to CTI summary statistics (post-hoc, weights not identifiable)")),
            ("Colour code", "Blue = hard-coded input / optimiser output; black = formula; red/green = failed/passed check.")]
    if exact is not None: rows.insert(3, ("Exact recovery", f"OLS R² = {exact['ols']['R2']:.6f}; NNLS R² = {exact['nnls']['R2']:.6f}"))
    if tgt is not None:
        best = min(fits, key=lambda x: fits[x][0][0]); rows.insert(3, ("Best-fitting form", f"{best} (loss {fits[best][0][0]:.3f}); see CTI_Fit and Weight_Spread."))
    for i, (a_, b_) in enumerate(rows, 3):
        put(ws, i, 1, a_, BO).alignment = WR; put(ws, i, 2, b_).alignment = WR; ws.row_dimensions[i].height = 36
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 110
    wb.save(path)

# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", required=True); ap.add_argument("--sheet")
    ap.add_argument("--metrics", nargs="+", required=True, help="metric column names")
    ap.add_argument("--invert", nargs="*", default=[], help="metrics where LOWER is better")
    ap.add_argument("--norm", choices=["minmax", "range100"], default="minmax")
    ap.add_argument("--cti-column"); ap.add_argument("--target"); ap.add_argument("--target-from")
    ap.add_argument("--models", help="JSON with keys cti, pmi (see DEFAULT_MODELS)")
    ap.add_argument("--seeds", type=int, default=6); ap.add_argument("--maxiter", type=int, default=250)
    ap.add_argument("--spread-tol", type=float, default=5.0); ap.add_argument("--out", default="CTI_audit.xlsx")
    ap.add_argument("--recalc", action="store_true", help="recalculate with LibreOffice if available")
    args = ap.parse_args()

    df = read_table(args.data, args.sheet)
    miss = [m for m in args.metrics + args.invert if m not in df.columns]
    if miss: sys.exit(f"Columns not found: {miss}\nAvailable: {list(df.columns)}")
    X = df[args.metrics].apply(pd.to_numeric, errors="coerce")
    ok = X.notna().all(axis=1)
    if args.cti_column: ok &= pd.to_numeric(df[args.cti_column], errors="coerce").notna()
    if (~ok).any(): print(f"[warn] dropping {(~ok).sum()} rows with missing/non-numeric values")
    X = X[ok]; n = len(X)
    simulated = any("SIMULATED" in str(v).upper() for col in df.columns if "status" in str(col).lower() for v in df[col])
    Z = normalise(X, set(args.invert), args.norm, (0, 100)); M = Z.values

    exact = None
    if args.cti_column:
        y = pd.to_numeric(df.loc[ok, args.cti_column]).values
        exact = exact_recovery(M, y); exact.update(_max=y.max(), _mean=y.mean(), _std=y.std(ddof=1))
        print(f"[exact] OLS R²={exact['ols']['R2']:.6f}  NNLS R²={exact['nnls']['R2']:.6f}")

    tgt = None
    if args.target: tgt = parse_target_kv(args.target)
    elif args.target_from: tgt = parse_target_from(args.target_from)
    elif exact is not None:
        tgt = None
    else: sys.exit("Provide --cti-column, --target, or --target-from.")

    fits, spread = {}, {}
    if tgt is not None:
        tv = np.array([tgt[k] for k in KEYS]); scale = 0.005
        for kind in ["simple", "geo", "affine"]:
            sols = fit(kind, M, tv, scale, args.seeds, args.maxiter); fits[kind] = [sols[0]]
            best = sols[0][0]; keep = [s for s in sols if s[0] <= max(2 * best, best + args.spread_tol)]
            W = np.array([s[1] for s in keep])
            spread[kind] = dict(lo=W.min(0), hi=W.max(0), mean=W.mean(0), kept=len(keep), seeds=args.seeds)
            print(f"[fit] {kind:7s} loss={best:9.3f}  w={np.round(sols[0][1], 3)}")
        eq = np.full(M.shape[1], 1 / M.shape[1]); fits["equal"] = [(0, eq, [], stats(M @ eq))]
        fits = {k: fits[k] for k in ["simple", "equal", "geo", "affine"]}
        if exact is not None:  # exact mode + target: Model_Consistency needs distribution stats from target
            pass
    models = json.load(open(args.models)) if args.models else DEFAULT_MODELS
    build(args.out, args, args.metrics, tgt, fits, spread, exact, models, n, simulated)
    if args.recalc and shutil.which("soffice"):
        tmp = Path(args.out).parent / "_recalc"; tmp.mkdir(exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), args.out], check=True, capture_output=True)
        shutil.move(str(tmp / Path(args.out).name), args.out); tmp.rmdir()
    print(f"[done] {args.out}")

if __name__ == "__main__":
    main()
