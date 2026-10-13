#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPRODUCIBLE PHARMACEUTICAL-METRICS DATA WORKFLOW
=================================================

Purpose
-------
This script creates an auditable analysis package for a 120-row table with five
metrics. It explicitly distinguishes:
  1) values observed in a source dataset,
  2) transformations derived from observed values, and
  3) simulated values generated only for software testing / sensitivity analysis.

IMPORTANT SCIENTIFIC-INTEGRITY NOTE
-----------------------------------
The five metrics in the user's workbook are:
  - AI Model Accuracy (%)
  - Batch Failure Rate (%)
  - Production Cost Reduction (%)
  - Regulatory Compliance Score
  - System Resilience Score

A public pharmaceutical-batch dataset may contain real batch/process measurements,
but it does NOT automatically contain these five exact metrics. This script does
not invent a mapping from unrelated source columns to these metrics and does not
label simulated rows as real observations.

If you use the optional simulation mode, output rows are labelled
"SIMULATED — NOT OBSERVED". To create a real-data analysis, provide a CSV/XLSX
file containing the actual measured values and a data dictionary that defines
each metric. The script will then validate and analyze those supplied columns.

Potential public-source starting points (verify access, license, and variable
definitions before analysis):
  - Scientific Data article: https://doi.org/10.1038/s41597-022-01203-x
  - Related Figshare collection: https://doi.org/10.6084/m9.figshare.c.5645578.v1
  - FDA Data Dashboard: https://datadashboard.fda.gov/

The links are starting points, not a claim that they contain all five metrics.
The script intentionally does not scrape websites or silently fabricate data.

Install
-------
    python -m pip install pandas numpy scipy openpyxl

Examples
--------
1) Analyze a source file with real measurements:
    python pharma_data_workflow.py --input measured_data.xlsx --sheet Sheet1 \
        --output real_data_analysis.xlsx --mode observed

2) Create a clearly labelled simulated test dataset using a target workbook:
    python pharma_data_workflow.py --target-workbook kham.xlsx \
        --output simulated_test_workbook.xlsx --mode simulate --n 120

3) Validate data only (no simulation):
    python pharma_data_workflow.py --input measured_data.csv \
        --output validated_real_data.xlsx --mode observed

Expected source columns for observed mode (rename with --columns if needed):
    AI_Model_Accuracy (%), Batch_Failure_Rate (%),
    Production_Cost_Reduction (%), Regulatory_Compliance_Score,
    System_Resilience_Score

For transparent reporting, retain the original data, source URL/citation, retrieval
date, license, data dictionary, preprocessing log, and any exclusions.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from scipy import stats

METRICS = [
    "AI_Model_Accuracy (%)",
    "Batch_Failure_Rate (%)",
    "Production_Cost_Reduction (%)",
    "Regulatory_Compliance_Score",
    "System_Resilience_Score",
]

PUBLIC_STARTING_POINTS = [
    {
        "name": "Scientific Data pharmaceutical batch dataset article",
        "url": "https://doi.org/10.1038/s41597-022-01203-x",
        "note": (
            "Potential source of real batch/process measurements. Verify exact "
            "columns and definitions; do not assume it contains all five metrics."
        ),
    },
    {
        "name": "Related Figshare collection",
        "url": "https://doi.org/10.6084/m9.figshare.c.5645578.v1",
        "note": (
            "Potential downloadable data collection associated with the article. "
            "Verify file metadata, license, and variable definitions."
        ),
    },
    {
        "name": "FDA Data Dashboard",
        "url": "https://datadashboard.fda.gov/",
        "note": (
            "Official regulatory data resource. Inspection/compliance data are not "
            "automatically equivalent to a numeric compliance score per batch."
        ),
    },
]


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def sha256_file(path: Path) -> str:
    """Return SHA-256 hash for reproducibility and file-integrity tracking."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_table(path: Path, sheet: Optional[str] = None) -> pd.DataFrame:
    """Read CSV, XLSX, or XLS input without modifying the source file."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=sheet or 0)
    raise ValueError(f"Unsupported input format: {suffix}. Use CSV or Excel.")


def normalize_column_name(name: object) -> str:
    """Normalize column labels conservatively for matching; preserve source data."""
    text = str(name).strip()
    return " ".join(text.replace("\n", " ").split())


def resolve_metric_columns(
    frame: pd.DataFrame,
    requested: Optional[Sequence[str]] = None,
) -> Dict[str, str]:
    """
    Resolve canonical metric names to actual source columns.
    If --columns is supplied, it must list five comma-separated names in METRICS order.
    """
    available = {normalize_column_name(c): c for c in frame.columns}
    names = list(requested) if requested else METRICS
    if len(names) != len(METRICS):
        raise ValueError(
            f"Exactly {len(METRICS)} source column names are required; got {len(names)}."
        )

    mapping: Dict[str, str] = {}
    missing = []
    for canonical, requested_name in zip(METRICS, names):
        key = normalize_column_name(requested_name)
        if key in available:
            mapping[canonical] = available[key]
        else:
            missing.append(f"{canonical} <- {requested_name}")
    if missing:
        raise ValueError(
            "Required metric columns were not found:\n  - "
            + "\n  - ".join(missing)
            + "\nUse --columns with the exact five source column names, in this order. "
              "Do not map unrelated variables merely to make the script run."
        )
    return mapping


def validate_observed_data(
    frame: pd.DataFrame,
    mapping: Dict[str, str],
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Validate source measurements. Invalid entries become NaN only in the analysis
    copy; the untouched source is preserved in the workbook.
    """
    cleaned = pd.DataFrame(index=frame.index)
    audit = []
    for metric, source_col in mapping.items():
        original = frame[source_col]
        numeric = pd.to_numeric(original, errors="coerce")
        invalid_text = int((original.notna() & numeric.isna()).sum())
        missing = int(numeric.isna().sum())
        cleaned[metric] = numeric
        audit.append({
            "Metric": metric,
            "Source column": str(source_col),
            "Rows in source": len(frame),
            "Missing or nonnumeric": missing,
            "Non-numeric values among nonmissing cells": invalid_text,
            "Observed minimum": numeric.min(),
            "Observed maximum": numeric.max(),
            "Observed mean": numeric.mean(),
            "Observed SD (sample)": numeric.std(ddof=1),
        })

    # Flag rather than silently delete suspicious values.
    bounds = {
        "AI_Model_Accuracy (%)": (0, 100),
        "Batch_Failure_Rate (%)": (0, 100),
        "Production_Cost_Reduction (%)": (None, None),
        "Regulatory_Compliance_Score": (None, None),
        "System_Resilience_Score": (None, None),
    }
    for metric, (lower, upper) in bounds.items():
        if metric not in cleaned:
            continue
        values = cleaned[metric]
        if lower is not None:
            count = int((values < lower).sum())
            if count:
                logging.warning("%s has %d values below %s.", metric, count, lower)
        if upper is not None:
            count = int((values > upper).sum())
            if count:
                logging.warning("%s has %d values above %s.", metric, count, upper)

    return cleaned, pd.DataFrame(audit)


def descriptive_statistics(frame: pd.DataFrame) -> pd.DataFrame:
    """Compute transparent descriptive statistics and t-based 95% CIs."""
    rows = []
    for col in frame.columns:
        x = pd.to_numeric(frame[col], errors="coerce").dropna().to_numpy(dtype=float)
        n = len(x)
        if n == 0:
            rows.append({"Variable": col, "N": 0})
            continue

        mean = float(np.mean(x))
        sd = float(np.std(x, ddof=1)) if n > 1 else np.nan
        se = sd / math.sqrt(n) if n > 1 else np.nan
        tcrit = float(stats.t.ppf(0.975, df=n - 1)) if n > 1 else np.nan
        rows.append({
            "Variable": col,
            "N": n,
            "Mean": mean,
            "SD (sample)": sd,
            "SE": se,
            "95% CI lower (mean)": mean - tcrit * se if n > 1 else np.nan,
            "95% CI upper (mean)": mean + tcrit * se if n > 1 else np.nan,
            "Minimum": float(np.min(x)),
            "Q1": float(np.quantile(x, 0.25)),
            "Median": float(np.median(x)),
            "Q3": float(np.quantile(x, 0.75)),
            "Maximum": float(np.max(x)),
            "Missing": int(frame[col].isna().sum()),
        })
    return pd.DataFrame(rows)


def simulate_from_target(
    target_path: Path,
    n: int,
    seed: int,
    sheet: Optional[str] = None,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Generate simulated test data based on target workbook's column means and SDs.
    This is NOT data collection and MUST NOT be reported as real-world observations.
    The method preserves target means/SDs approximately, subject to range constraints.
    """
    target_raw = read_table(target_path, sheet)
    mapping = resolve_metric_columns(target_raw)
    target, audit = validate_observed_data(target_raw, mapping)

    target_means = target.mean()
    target_sds = target.std(ddof=1)
    rng = np.random.default_rng(seed)
    out = pd.DataFrame(index=range(n))
    diagnostics = []

    for metric in METRICS:
        observed = target[metric].dropna()
        if len(observed) < 2:
            raise ValueError(f"At least two valid values are needed for {metric}.")
        mean = float(target_means[metric])
        sd = float(target_sds[metric])
        low = float(observed.min())
        high = float(observed.max())

        z = rng.normal(size=n)
        z_sd = float(np.std(z, ddof=1))
        z = (z - float(np.mean(z))) / z_sd
        values = mean + z * sd

        # Keep generated values within the target workbook's observed range.
        # Clipping may alter mean/SD; report the resulting difference explicitly.
        values = np.clip(values, low, high)
        out[metric] = values
        diagnostics.append({
            "Metric": metric,
            "Target mean": mean,
            "Target sample SD": sd,
            "Target min": low,
            "Target max": high,
            "Generated mean": float(np.mean(values)),
            "Generated sample SD": float(np.std(values, ddof=1)),
            "Mean difference": float(np.mean(values) - mean),
            "Generation method": "Seeded normal draws, standardized then clipped",
            "Status": "SIMULATED — NOT OBSERVED",
        })

    out.insert(0, "Record_ID", [f"SIM-{i:03d}" for i in range(1, n + 1)])
    out.insert(1, "Data_Status", "SIMULATED — NOT OBSERVED")
    return out, pd.DataFrame(diagnostics), audit


def make_sources_sheet() -> pd.DataFrame:
    return pd.DataFrame(PUBLIC_STARTING_POINTS)


def write_workbook(
    output: Path,
    data: pd.DataFrame,
    summary: pd.DataFrame,
    audit: pd.DataFrame,
    sources: pd.DataFrame,
    metadata: dict,
    raw_source: Optional[pd.DataFrame] = None,
) -> None:
    """Write analysis output, provenance, validation, and source references."""
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="Analyzed_Data")
        summary.to_excel(writer, index=False, sheet_name="Descriptive_Stats")
        audit.to_excel(writer, index=False, sheet_name="Validation_Audit")
        sources.to_excel(writer, index=False, sheet_name="Source_Registry")
        pd.DataFrame([{"Key": k, "Value": json.dumps(v, ensure_ascii=False)
                       if isinstance(v, (dict, list)) else str(v)}
                      for k, v in metadata.items()]).to_excel(
            writer, index=False, sheet_name="Provenance"
        )
        if raw_source is not None:
            raw_source.to_excel(writer, index=False, sheet_name="Original_Source")

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = load_workbook(output)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col_cells in ws.columns:
                width = min(max(max(
                    len(str(c.value)) if c.value is not None else 0
                    for c in col_cells[:min(ws.max_row, 150) + 1]
                ) + 2, 12), 55)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width
        wb.save(output)
    except Exception as exc:
        logging.warning("Workbook styling skipped: %s", exc)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Auditable workflow for observed or clearly labelled simulated data."
    )
    parser.add_argument("--mode", choices=["observed", "simulate"], required=True)
    parser.add_argument("--input", type=Path,
                        help="Observed CSV/XLSX file for --mode observed.")
    parser.add_argument("--target-workbook", type=Path,
                        help="Workbook with five target columns for --mode simulate.")
    parser.add_argument("--sheet", default=None, help="Optional Excel sheet name.")
    parser.add_argument("--output", type=Path, required=True,
                        help="Output .xlsx path.")
    parser.add_argument("--n", type=int, default=120,
                        help="Number of simulated rows (default: 120).")
    parser.add_argument("--seed", type=int, default=20261013,
                        help="Random seed for reproducibility.")
    parser.add_argument("--columns", default=None,
                        help="Five exact source column names separated by commas, in metric order.")
    args = parser.parse_args()

    setup_logging()
    if args.output.suffix.lower() != ".xlsx":
        parser.error("--output must end in .xlsx")

    if args.mode == "observed":
        if not args.input:
            parser.error("--input is required for --mode observed")
        if not args.input.exists():
            parser.error(f"Input file not found: {args.input}")
        raw = read_table(args.input, args.sheet)
        requested = [s.strip() for s in args.columns.split(",")] if args.columns else None
        mapping = resolve_metric_columns(raw, requested)
        cleaned, audit = validate_observed_data(raw, mapping)
        summary = descriptive_statistics(cleaned)
        data = cleaned.copy()
        data.insert(0, "Record_ID", [f"OBS-{i:04d}" for i in range(1, len(data) + 1)])
        data.insert(1, "Data_Status", "OBSERVED IN USER-SUPPLIED SOURCE; provenance must be documented")
        metadata = {
            "mode": "observed",
            "created_utc": datetime.now(timezone.utc).isoformat(),
            "input_file": str(args.input.resolve()),
            "input_sha256": sha256_file(args.input),
            "source_rows": len(raw),
            "source_columns": list(map(str, raw.columns)),
            "note": "Observed means values found in the supplied file; external provenance is not independently verified by this script.",
        }
        write_workbook(args.output, data, summary, audit, make_sources_sheet(), metadata, raw)
        logging.info("Analyzed observed input. No rows were synthetically generated.")
    else:
        if not args.target_workbook:
            parser.error("--target-workbook is required for --mode simulate")
        if not args.target_workbook.exists():
            parser.error(f"Target workbook not found: {args.target_workbook}")
        if args.n < 2:
            parser.error("--n must be >= 2")
        data, diagnostics, audit = simulate_from_target(
            args.target_workbook, args.n, args.seed, args.sheet
        )
        metric_frame = data[METRICS]
        summary = descriptive_statistics(metric_frame)
        metadata = {
            "mode": "simulate",
            "created_utc": datetime.now(timezone.utc).isoformat(),
            "target_workbook": str(args.target_workbook.resolve()),
            "target_sha256": sha256_file(args.target_workbook),
            "n": args.n,
            "seed": args.seed,
            "note": "All generated rows are simulated and must not be described as observed real-world data.",
            "generation_diagnostics": diagnostics.to_dict(orient="records"),
        }
        write_workbook(args.output, data, summary, audit, make_sources_sheet(), metadata)
        logging.warning("SIMULATED data written. Do not report it as real observed data.")

    logging.info("Output saved to %s", args.output.resolve())
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        logging.error("%s", exc)
        raise SystemExit(1)
